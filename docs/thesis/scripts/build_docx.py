from __future__ import annotations

import argparse
import re
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[3]
THESIS = ROOT / "docs" / "thesis"
SOURCE = THESIS / "manuscript.md"
DEFAULT_OUTPUT = THESIS / "UNIBUS_졸업논문_권재원외5인.docx"

FONT_KO = "Batang"
FONT_SANS = "Malgun Gothic"
FONT_EN = "Times New Roman"
NAVY = "173873"
BLUE = "DCEBFA"
LIGHT = "F4F6F9"
LINE = "AAB4C3"
INK = "111827"
MUTED = "596579"


@dataclass
class HeadingRecord:
    line_index: int
    level: int
    text: str
    bookmark: str


@dataclass
class CaptionRecord:
    line_index: int
    label: str
    title: str
    bookmark: str
    kind: str


def set_run_font(run, name: str, size: float, bold: bool | None = None, italic: bool | None = None, color: str = INK) -> None:
    run.font.name = name
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def set_style_font(style, name: str, size: float, bold: bool = False, color: str = INK) -> None:
    style.font.name = name
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor.from_string(color)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style._element.rPr.rFonts.set(qn("w:ascii"), name)
    style._element.rPr.rFonts.set(qn("w:hAnsi"), name)


def set_page_geometry(section) -> None:
    section.page_width = Mm(190)
    section.page_height = Mm(260)
    section.top_margin = Mm(15)
    section.bottom_margin = Mm(15)
    section.left_margin = Mm(25)
    section.right_margin = Mm(25)
    section.header_distance = Mm(8)
    section.footer_distance = Mm(8)


def set_update_fields(document: Document) -> None:
    settings = document.settings._element
    update = settings.find(qn("w:updateFields"))
    if update is None:
        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(qn("w:val"), "true")


def set_page_number_format(section, fmt: str, start: int) -> None:
    sect_pr = section._sectPr
    pg_num = sect_pr.find(qn("w:pgNumType"))
    if pg_num is None:
        pg_num = OxmlElement("w:pgNumType")
        sect_pr.append(pg_num)
    pg_num.set(qn("w:fmt"), fmt)
    pg_num.set(qn("w:start"), str(start))


def append_field(paragraph, instruction: str, placeholder: str = "0") -> None:
    begin_run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    begin_run._r.append(begin)

    instr_run = paragraph.add_run()
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = instruction
    instr_run._r.append(instr)

    separate_run = paragraph.add_run()
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    separate_run._r.append(separate)

    value_run = paragraph.add_run(placeholder)
    set_run_font(value_run, FONT_KO, 9)

    end_run = paragraph.add_run()
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    end_run._r.append(end)


def add_bookmark(paragraph, name: str, bookmark_id: int) -> None:
    start = OxmlElement("w:bookmarkStart")
    start.set(qn("w:id"), str(bookmark_id))
    start.set(qn("w:name"), name)
    end = OxmlElement("w:bookmarkEnd")
    end.set(qn("w:id"), str(bookmark_id))
    paragraph._p.insert(0, start)
    paragraph._p.append(end)


def add_page_number_footer(section, font_name: str = FONT_KO) -> None:
    section.footer.is_linked_to_previous = False
    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    append_field(paragraph, " PAGE ", "1")
    for run in paragraph.runs:
        set_run_font(run, font_name, 9, color=MUTED)


def add_body_header(section) -> None:
    section.header.is_linked_to_previous = False
    paragraph = section.header.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run("UNIBUS  |  2D/3D 대학 셔틀버스 통합 운행 정보 시스템")
    set_run_font(run, FONT_SANS, 7.5, color=MUTED)
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), "D4DAE3")
    border.append(bottom)
    p_pr.append(border)


def add_centered_paragraph(document: Document, text: str, size: float, bold: bool = False, color: str = INK, space_after: float = 0) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.first_line_indent = Pt(0)
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    paragraph.paragraph_format.space_after = Pt(space_after)
    run = paragraph.add_run(text)
    set_run_font(run, FONT_KO, size, bold=bold, color=color)


def vertical_space(document: Document, points: float) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    paragraph.paragraph_format.line_spacing = Pt(1)
    paragraph.paragraph_format.space_after = Pt(points)


def add_cover(document: Document) -> None:
    vertical_space(document, 8)
    add_centered_paragraph(document, "정보통신공학과 졸업논문", 15, bold=True, color=NAVY, space_after=18)
    add_centered_paragraph(document, "2D/3D 디지털 트윈 기반", 22, bold=True, space_after=4)
    add_centered_paragraph(document, "대학 셔틀버스 통합 운행", 22, bold=True, space_after=4)
    add_centered_paragraph(document, "정보 시스템의 설계 및 구현", 22, bold=True, space_after=10)
    add_centered_paragraph(document, "순천향대학교 프로토타입", 14, bold=True, color=NAVY, space_after=8)
    add_centered_paragraph(document, "Design and Implementation of an Integrated University Shuttle Bus", 11, color=MUTED)
    add_centered_paragraph(document, "Information System Based on a 2D/3D Digital Twin", 11, color=MUTED)
    vertical_space(document, 18)
    add_centered_paragraph(document, "지도교수  박 동 규", 12, space_after=15)
    add_centered_paragraph(document, "20214157  권 재 원", 12, bold=True, space_after=5)
    add_centered_paragraph(document, "박 소 정   ·   손 성 윤   ·   김 동 하", 11.5, space_after=5)
    add_centered_paragraph(document, "정 영 선   ·   강 민 영", 11.5, space_after=15)
    add_centered_paragraph(document, "2026년 7월", 12, space_after=10)
    add_centered_paragraph(document, "순 천 향 대 학 교", 18, bold=True, color=NAVY, space_after=4)
    add_centered_paragraph(document, "정보통신공학과", 13, bold=True)


def add_inner_title(document: Document) -> None:
    document.add_page_break()
    vertical_space(document, 42)
    add_centered_paragraph(document, "2D/3D 디지털 트윈 기반 대학 셔틀버스", 19, bold=True, space_after=6)
    add_centered_paragraph(document, "통합 운행 정보 시스템의 설계 및 구현", 19, bold=True, space_after=13)
    add_centered_paragraph(document, "순천향대학교 프로토타입", 13, bold=True, color=NAVY, space_after=38)
    add_centered_paragraph(document, "이 논문을 정보통신공학과 졸업논문으로 제출함", 11.5, space_after=34)
    add_centered_paragraph(document, "2026년 7월", 11.5, space_after=28)
    add_centered_paragraph(document, "20214157  권 재 원", 12, bold=True, space_after=5)
    add_centered_paragraph(document, "박 소 정   ·   손 성 윤   ·   김 동 하", 11.5, space_after=5)
    add_centered_paragraph(document, "정 영 선   ·   강 민 영", 11.5, space_after=36)
    add_centered_paragraph(document, "순천향대학교 정보통신공학과", 14, bold=True, color=NAVY)


def add_approval_page(document: Document) -> None:
    document.add_page_break()
    vertical_space(document, 48)
    add_centered_paragraph(document, "졸업논문 승인", 20, bold=True, color=NAVY, space_after=44)
    add_centered_paragraph(document, "권재원 외 5인의 졸업논문을 심사하고", 12, space_after=5)
    add_centered_paragraph(document, "정보통신공학과 졸업논문으로 인정함", 12, space_after=48)
    add_centered_paragraph(document, "지도교수   박 동 규   (인)", 12, space_after=66)
    add_centered_paragraph(document, "2026년 7월", 11.5, space_after=36)
    add_centered_paragraph(document, "순 천 향 대 학 교", 18, bold=True, color=NAVY, space_after=4)
    add_centered_paragraph(document, "정보통신공학과", 13, bold=True)


def configure_styles(document: Document) -> None:
    styles = document.styles
    normal = styles["Normal"]
    set_style_font(normal, FONT_KO, 11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = 1.8
    normal.paragraph_format.first_line_indent = Pt(20)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.widow_control = True

    for name, size, before, after, align in [
        ("Heading 1", 16, 18, 13, WD_ALIGN_PARAGRAPH.CENTER),
        ("Heading 2", 13, 14, 7, WD_ALIGN_PARAGRAPH.LEFT),
        ("Heading 3", 11.5, 11, 5, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        style = styles[name]
        set_style_font(style, FONT_KO, size, bold=True, color=NAVY if name == "Heading 1" else INK)
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.keep_together = True

    for name, size, bold, align in [
        ("Thesis Front Title", 16, True, WD_ALIGN_PARAGRAPH.CENTER),
        ("Thesis Figure Caption", 9.5, False, WD_ALIGN_PARAGRAPH.CENTER),
        ("Thesis Table Caption", 9.5, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Thesis List Entry", 10.5, False, WD_ALIGN_PARAGRAPH.LEFT),
        ("Thesis English", 10.5, False, WD_ALIGN_PARAGRAPH.JUSTIFY),
        ("Thesis Reference", 9.5, False, WD_ALIGN_PARAGRAPH.LEFT),
    ]:
        if name not in styles:
            styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        style = styles[name]
        set_style_font(style, FONT_EN if name == "Thesis English" else FONT_KO, size, bold=bold, color=INK)
        style.paragraph_format.alignment = align
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_after = Pt(0)
        style.paragraph_format.line_spacing = 1.5 if name == "Thesis English" else 1.35
        if name == "Thesis List Entry":
            style.paragraph_format.line_spacing = 1.1
        if "Caption" in name:
            style.paragraph_format.keep_with_next = name == "Thesis Table Caption"
            style.paragraph_format.keep_together = True
            style.paragraph_format.space_before = Pt(5)
            style.paragraph_format.space_after = Pt(5)
        if name == "Thesis Reference":
            style.paragraph_format.left_indent = Pt(20)
            style.paragraph_format.first_line_indent = Pt(-20)
            style.paragraph_format.space_after = Pt(4)


def preparse(lines: list[str]) -> tuple[list[HeadingRecord], list[CaptionRecord], list[CaptionRecord]]:
    headings: list[HeadingRecord] = []
    figures: list[CaptionRecord] = []
    tables: list[CaptionRecord] = []
    h_count = f_count = t_count = 0
    for idx, line in enumerate(lines):
        heading = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading:
            h_count += 1
            headings.append(HeadingRecord(idx, len(heading.group(1)), heading.group(2).strip(), f"heading_{h_count:03d}"))
            continue
        fig = re.match(r"^\[FIGURE:([^|]+)\|([^|]+)\|([^|]+)\|([0-9.]+)\]$", line)
        if fig:
            f_count += 1
            figures.append(CaptionRecord(idx, fig.group(2), fig.group(3), f"figure_{f_count:03d}", "figure"))
            continue
        table = re.match(r"^\[TABLE:([^|]+)\|([^]]+)\]$", line)
        if table:
            t_count += 1
            tables.append(CaptionRecord(idx, table.group(1), table.group(2), f"table_{t_count:03d}", "table"))
    return headings, figures, tables


def add_front_title(document: Document, text: str) -> None:
    paragraph = document.add_paragraph(style="Thesis Front Title")
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(18)
    run = paragraph.add_run(text)
    set_run_font(run, FONT_KO, 16, bold=True, color=NAVY)


def add_pageref_entry(document: Document, text: str, bookmark: str, level: int = 1) -> None:
    paragraph = document.add_paragraph(style="Thesis List Entry")
    paragraph.paragraph_format.left_indent = Pt((level - 1) * 16)
    paragraph.paragraph_format.space_after = Pt(1 if level == 1 else 0)
    tab_stops = paragraph.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(13.6 - (level - 1) * 0.2), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    run = paragraph.add_run(text)
    set_run_font(run, FONT_KO, 9.6 if level > 1 else 10.1, bold=level == 1)
    paragraph.add_run("\t")
    append_field(paragraph, f" PAGEREF {bookmark} \\h ", "0")


def add_contents_pages(document: Document, headings: list[HeadingRecord], figures: list[CaptionRecord], tables: list[CaptionRecord]) -> None:
    add_front_title(document, "목 차")
    for record in headings:
        if record.text == "국문초록":
            add_pageref_entry(document, "국문초록", record.bookmark, 1)
        elif record.text.startswith("제") or record.text in {"참고문헌", "Abstract"} or record.text.startswith("부록"):
            add_pageref_entry(document, record.text, record.bookmark, 1)
        elif record.level == 2 and not record.text.startswith("Design and Implementation"):
            add_pageref_entry(document, record.text, record.bookmark, 2)

    document.add_page_break()
    add_front_title(document, "표 목 차")
    for record in tables:
        add_pageref_entry(document, f"{record.label}. {record.title}", record.bookmark, 1)

    document.add_page_break()
    add_front_title(document, "그 림 목 차")
    for record in figures:
        add_pageref_entry(document, f"{record.label}. {record.title}", record.bookmark, 1)
    document.add_page_break()


def add_inline_markdown(paragraph, text: str, english: bool = False) -> None:
    pattern = re.compile(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)")
    cursor = 0
    for match in pattern.finditer(text):
        if match.start() > cursor:
            run = paragraph.add_run(text[cursor:match.start()])
            set_run_font(run, FONT_EN if english else FONT_KO, 10.5 if english else 11)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, FONT_EN if english else FONT_KO, 10.5 if english else 11, bold=True)
        elif token.startswith("*"):
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, FONT_EN if english else FONT_KO, 10.5 if english else 11, italic=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, "Menlo", 8.5, color=NAVY)
        cursor = match.end()
    if cursor < len(text):
        run = paragraph.add_run(text[cursor:])
        set_run_font(run, FONT_EN if english else FONT_KO, 10.5 if english else 11)


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top: int = 90, start: int = 100, bottom: int = 90, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color: str = LINE, size: int = 5) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            borders.append(node)
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_cm: list[float]) -> None:
    total_twips = int(sum(widths_cm) * 567)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total_twips))
    tbl_w.set(qn("w:type"), "dxa")
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_cm:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(int(width * 567)))
        grid.append(col)

    for row in table.rows:
        for cell, width in zip(row.cells, widths_cm):
            cell.width = Cm(width)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(int(width * 567)))
            tc_w.set(qn("w:type"), "dxa")


def calculate_widths(rows: list[list[str]]) -> list[float]:
    columns = len(rows[0])
    scores = []
    for col in range(columns):
        longest = max(len(row[col]) for row in rows)
        average = sum(len(row[col]) for row in rows) / len(rows)
        scores.append(max(5, min(35, longest * 0.55 + average * 0.45)))
    minimum = 1.35 if columns >= 6 else 1.7
    raw = [max(minimum, score) for score in scores]
    scale = 13.8 / sum(raw)
    widths = [max(minimum, value * scale) for value in raw]
    if sum(widths) > 13.8:
        excess = sum(widths) - 13.8
        adjustable = [max(0, width - minimum) for width in widths]
        total_adjustable = sum(adjustable) or 1
        widths = [width - excess * adj / total_adjustable for width, adj in zip(widths, adjustable)]
    return widths


def add_table(document: Document, rows: list[list[str]], caption: CaptionRecord, bookmark_id: int) -> None:
    caption_paragraph = document.add_paragraph(style="Thesis Table Caption")
    caption_paragraph.add_run(f"{caption.label}. {caption.title}")
    add_bookmark(caption_paragraph, caption.bookmark, bookmark_id)
    for run in caption_paragraph.runs:
        set_run_font(run, FONT_KO, 9.5, bold=True)

    table = document.add_table(rows=len(rows), cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    widths = calculate_widths(rows)
    set_table_geometry(table, widths)
    set_table_borders(table)

    for row_idx, data_row in enumerate(rows):
        row = table.rows[row_idx]
        set_row_cant_split(row)
        if row_idx == 0:
            set_repeat_table_header(row)
        for col_idx, value in enumerate(data_row):
            cell = row.cells[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            set_cell_shading(cell, NAVY if row_idx == 0 else (LIGHT if row_idx % 2 == 0 else "FFFFFF"))
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER if (row_idx == 0 or len(value) < 18) else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.15
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, FONT_KO, 8.3 if len(rows[0]) >= 6 else 8.8, bold=row_idx == 0, color="FFFFFF" if row_idx == 0 else INK)
    after = document.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_figure(document: Document, marker: re.Match[str], caption: CaptionRecord, bookmark_id: int) -> None:
    relative_path = marker.group(1)
    width = float(marker.group(4))
    image_path = THESIS / relative_path
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run()
    inline = run.add_picture(str(image_path), width=Cm(width))
    doc_pr = inline._inline.docPr
    doc_pr.set("descr", caption.title)

    cap = document.add_paragraph(style="Thesis Figure Caption")
    cap_run = cap.add_run(f"{caption.label}. {caption.title}")
    set_run_font(cap_run, FONT_KO, 9.5, bold=True)
    add_bookmark(cap, caption.bookmark, bookmark_id)


def parse_table_rows(lines: list[str], start: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    idx = start
    while idx < len(lines) and lines[idx].startswith("|"):
        cells = [cell.strip() for cell in lines[idx].strip().strip("|").split("|")]
        if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
            rows.append(cells)
        idx += 1
    return rows, idx


def render_lines(
    document: Document,
    lines: list[str],
    start: int,
    end: int,
    heading_map: dict[int, HeadingRecord],
    figure_map: dict[int, CaptionRecord],
    table_map: dict[int, CaptionRecord],
    bookmark_ids: dict[str, int],
) -> None:
    idx = start
    in_references = False
    in_english = False
    force_page_break = False
    while idx < end:
        line = lines[idx].strip()
        if not line:
            idx += 1
            continue
        if line == "[[PAGEBREAK]]":
            force_page_break = True
            idx += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.+)$", line)
        if heading_match:
            record = heading_map[idx]
            text = record.text
            level = record.level
            if text == "국문초록":
                paragraph = document.add_paragraph(style="Thesis Front Title")
            elif text == "Abstract":
                paragraph = document.add_paragraph(style="Heading 1")
                in_english = True
                in_references = False
            elif in_english and level == 2 and text.startswith("Design and Implementation"):
                paragraph = document.add_paragraph(style="Thesis Front Title")
                paragraph.paragraph_format.space_after = Pt(18)
            else:
                paragraph = document.add_paragraph(style=f"Heading {min(level, 3)}")
            if force_page_break:
                paragraph.paragraph_format.page_break_before = True
                force_page_break = False
            paragraph.add_run(text)
            add_bookmark(paragraph, record.bookmark, bookmark_ids[record.bookmark])
            for run in paragraph.runs:
                set_run_font(run, FONT_EN if in_english else FONT_KO, 14 if in_english and level == 2 else (16 if level == 1 else 13), bold=True, color=NAVY if level == 1 else INK)
            in_references = text == "참고문헌"
            if text.startswith("부록"):
                in_references = False
            idx += 1
            continue

        table_marker = re.match(r"^\[TABLE:([^|]+)\|([^]]+)\]$", line)
        if table_marker:
            rows, next_idx = parse_table_rows(lines, idx + 1)
            record = table_map[idx]
            add_table(document, rows, record, bookmark_ids[record.bookmark])
            idx = next_idx
            continue

        figure_marker = re.match(r"^\[FIGURE:([^|]+)\|([^|]+)\|([^|]+)\|([0-9.]+)\]$", line)
        if figure_marker:
            record = figure_map[idx]
            add_figure(document, figure_marker, record, bookmark_ids[record.bookmark])
            idx += 1
            continue

        style = "Thesis Reference" if in_references else ("Thesis English" if in_english else "Normal")
        paragraph = document.add_paragraph(style=style)
        if line.startswith("**주요어:**") or line.startswith("**Keywords:**"):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(10)
        if line.startswith("**RQ"):
            paragraph.paragraph_format.first_line_indent = Pt(0)
            paragraph.paragraph_format.space_before = Pt(4)
        add_inline_markdown(paragraph, line, english=in_english or in_references)
        idx += 1


def build(output: Path) -> None:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    headings, figures, tables = preparse(lines)
    heading_map = {record.line_index: record for record in headings}
    figure_map = {record.line_index: record for record in figures}
    table_map = {record.line_index: record for record in tables}
    all_bookmarks = [record.bookmark for record in headings + figures + tables]
    bookmark_ids = {name: idx + 1 for idx, name in enumerate(all_bookmarks)}

    document = Document()
    configure_styles(document)
    set_update_fields(document)
    first = document.sections[0]
    set_page_geometry(first)
    first.header.is_linked_to_previous = False
    first.footer.is_linked_to_previous = False

    add_cover(document)
    add_inner_title(document)
    add_approval_page(document)

    front_section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(front_section)
    set_page_number_format(front_section, "lowerRoman", 1)
    add_page_number_footer(front_section)
    front_section.header.is_linked_to_previous = False
    front_section.header.paragraphs[0].clear()
    add_contents_pages(document, headings, figures, tables)

    first_chapter_index = next(record.line_index for record in headings if record.text.startswith("제1장"))
    render_lines(document, lines, 0, first_chapter_index, heading_map, figure_map, table_map, bookmark_ids)

    body_section = document.add_section(WD_SECTION.NEW_PAGE)
    set_page_geometry(body_section)
    set_page_number_format(body_section, "decimal", 1)
    add_page_number_footer(body_section)
    add_body_header(body_section)
    render_lines(document, lines, first_chapter_index, len(lines), heading_map, figure_map, table_map, bookmark_ids)

    core = document.core_properties
    core.title = "2D/3D 디지털 트윈 기반 대학 셔틀버스 통합 운행 정보 시스템의 설계 및 구현"
    core.subject = "순천향대학교 정보통신공학과 졸업논문"
    core.author = "권재원, 박소정, 손성윤, 김동하, 정영선, 강민영"
    core.keywords = "대학 셔틀버스, 디지털 트윈, PWA, 실시간 위치, Supabase"
    core.comments = "2026-07-15 기준 UNIBUS 프로토타입 연구"

    output.parent.mkdir(parents=True, exist_ok=True)
    document.save(output)
    print(f"Saved {output}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output", type=Path, default=DEFAULT_OUTPUT)
    args = parser.parse_args()
    build(args.output.resolve())


if __name__ == "__main__":
    main()
